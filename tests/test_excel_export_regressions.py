import io
import json
from contextlib import closing
from pathlib import Path
import subprocess
import sys
import textwrap
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
import pytest

from astrbot_plugin_office_assistant.domain.workbook.contracts import (
    CreateWorkbookRequest,
    ExportWorkbookRequest,
    WriteRowsOptions,
    WriteRowsRequest,
)
from astrbot_plugin_office_assistant.domain.workbook.session_store import (
    WorkbookSessionStore,
)
from astrbot_plugin_office_assistant.services.excel_script_templates import (
    build_runner_script,
)


def _run_export(
    tmp_path: Path,
    script: str,
    *,
    input_files: list[Path] | None = None,
    suffix: str = ".xlsx",
) -> tuple[Path, Path]:
    exec_dir = tmp_path / "exec"
    exec_dir.mkdir()
    output_path = tmp_path / f"output{suffix}"
    result_path = tmp_path / "result.json"
    runner_path = tmp_path / "runner.py"
    runner_path.write_text(
        build_runner_script(
            script=textwrap.dedent(script),
            exec_dir=str(exec_dir),
            input_files=[str(path) for path in input_files or []],
            output_path=str(output_path),
            result_path=str(result_path),
        ),
        encoding="utf-8",
    )
    completed = subprocess.run(
        [sys.executable, str(runner_path)],
        cwd=tmp_path,
        capture_output=True,
        text=True,
        timeout=30,
    )
    assert completed.returncode == 0, completed.stderr
    result = json.loads(result_path.read_text(encoding="utf-8"))
    assert result["success"] is True, result
    assert result["mode"] == "file"
    return output_path, exec_dir


@pytest.fixture
def macro_workbook(tmp_path: Path) -> tuple[Path, bytes]:
    # An inert payload tests OOXML archive retention without executing any macro.
    macro_payload = b"office assistant inert VBA preservation fixture"
    base_archive = io.BytesIO()
    workbook = Workbook()
    workbook.active.append(["Course", "Room"])
    workbook.active.append(["Physics", "101"])
    workbook.save(base_archive)
    with ZipFile(base_archive, "a") as archive:
        archive.writestr("xl/vbaProject.bin", macro_payload)
    workbook.vba_archive = ZipFile(io.BytesIO(base_archive.getvalue()))
    input_path = tmp_path / "input.xlsm"
    try:
        workbook.save(input_path)
    finally:
        workbook.vba_archive.close()
        workbook.close()
    return input_path, macro_payload


@pytest.mark.parametrize("save_kind", ["helper", "direct", "copy"])
def test_excel_runner_preserves_xlsm_vba(
    tmp_path: Path, macro_workbook: tuple[Path, bytes], save_kind: str
):
    input_path, macro_payload = macro_workbook
    scripts = {
        "helper": (
            "workbook = load_input_workbook(keep_vba=True)\n"
            "save_output_workbook(workbook)"
        ),
        "direct": (
            "workbook = load_input_workbook(keep_vba=True)\nworkbook.save(output_path)"
        ),
        "copy": "import shutil\nshutil.copyfile(input_files[0], output_path)",
    }
    output_path, _ = _run_export(
        tmp_path, scripts[save_kind], input_files=[input_path], suffix=".xlsm"
    )
    with ZipFile(output_path) as archive:
        assert archive.read("xl/vbaProject.bin") == macro_payload
        assert b"macroEnabled" in archive.read("[Content_Types].xml")
    with (
        closing(load_workbook(output_path, keep_vba=True)) as workbook,
        workbook.vba_archive,
    ):
        assert workbook.active["A2"].value == "Physics"
        assert workbook.calculation.forceFullCalc is True


_FORMAT_COUNTER = """
original_format = auto_format_workbook
def counting_format(workbook, **kwargs):
    with Path("format_calls.txt").open("a", encoding="utf-8") as counter:
        counter.write("formatted\\n")
    return original_format(workbook, **kwargs)
auto_format_workbook = counting_format
workbook = Workbook()
workbook.active.append(["Course", "Room"])
workbook.active.append(["Physics", "101"])
"""


@pytest.mark.parametrize("save_kind", ["helper", "direct", "replaced"])
def test_excel_runner_formats_only_new_saved_content(tmp_path: Path, save_kind: str):
    save_statement = (
        "workbook.save(output_path)"
        if save_kind == "direct"
        else "save_output_workbook(workbook)"
    )
    script = _FORMAT_COUNTER + save_statement + "\n"
    replaced = save_kind == "replaced"
    if replaced:
        script += """
replacement = Workbook()
replacement.active.append(["Course", "Room", "Teacher"])
replacement.active.append(["History", "201", "New teacher"])
openpyxl.writer.excel.save_workbook(replacement, output_path)
"""
    output_path, exec_dir = _run_export(tmp_path, script)
    expected_calls = 2 if replaced else 1
    assert (exec_dir / "format_calls.txt").read_text().splitlines() == [
        "formatted"
    ] * expected_calls
    with closing(load_workbook(output_path)) as workbook:
        assert workbook.active["A2"].value == ("History" if replaced else "Physics")
        assert workbook.active.freeze_panes == "A2"
        assert workbook.active.auto_filter.ref == ("A1:C2" if replaced else "A1:B2")


def test_structured_export_autofilter_starts_at_table_header(tmp_path: Path):
    store = WorkbookSessionStore(workspace_dir=tmp_path)
    workbook = store.create_workbook(CreateWorkbookRequest(filename="report.xlsx"))
    store.write_rows(
        WriteRowsRequest(
            workbook_id=workbook.workbook_id,
            sheet="Report",
            rows=[["Monthly report"]],
        )
    )
    store.write_rows(
        WriteRowsRequest(
            workbook_id=workbook.workbook_id,
            sheet="Report",
            start_row=3,
            rows=[["Name", "Amount"], ["Alice", 100]],
            options=WriteRowsOptions(autofilter=True),
        )
    )
    _, output_path = store.export_workbook(
        ExportWorkbookRequest(workbook_id=workbook.workbook_id)
    )
    with closing(load_workbook(output_path)) as exported:
        assert exported.active.auto_filter.ref == "A3:B4"
        assert exported.active["A1"].value == "Monthly report"
        assert exported.active["A3"].font.bold is True
