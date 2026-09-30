from concurrent.futures import ThreadPoolExecutor
from contextlib import nullcontext
from pathlib import Path
from threading import Barrier, Lock
from types import SimpleNamespace
import sys

import pytest

from astrbot_plugin_office_assistant.services import pdf_converter as converter_module
from astrbot_plugin_office_assistant.services.pdf_converter import PDFConverter


_BACKENDS = ["docx2pdf", "win32com", "libreoffice", "word", "tabula", "pdfplumber"]


def _fake_backend(monkeypatch, tmp_path: Path, backend: str, write_output):
    converter = PDFConverter.__new__(PDFConverter)
    converter.data_path = tmp_path
    converter._libreoffice_path = "soffice"
    office_backend = backend in {"docx2pdf", "win32com", "libreoffice"}
    input_path = tmp_path / ("report.docx" if office_backend else "report.pdf")
    input_path.write_bytes(b"source must remain unchanged")
    extension = ".pdf" if office_backend else ".docx" if backend == "word" else ".xlsx"

    class FakeDocument:
        def SaveAs(self, path, *args, **kwargs):
            write_output(Path(path))

        def ExportAsFixedFormat(self, kind, path):
            write_output(Path(path))

        def Close(self, **kwargs):
            pass

    app = SimpleNamespace(
        **{
            name: SimpleNamespace(Open=lambda *args, **kwargs: FakeDocument())
            for name in ("Documents", "Workbooks", "Presentations")
        }
    )
    monkeypatch.setattr(converter_module, "com_application", lambda _: nullcontext(app))
    monkeypatch.setattr(converter_module, "_WIN32COM_AVAILABLE", True)
    monkeypatch.setitem(
        sys.modules,
        "docx2pdf",
        SimpleNamespace(convert=lambda source, output: write_output(Path(output))),
    )
    monkeypatch.setitem(
        sys.modules,
        "pythoncom",
        SimpleNamespace(CoInitialize=lambda: None, CoUninitialize=lambda: None),
    )

    def fake_subprocess_run(command, **kwargs):
        output_dir = Path(command[command.index("--outdir") + 1])
        write_output(output_dir / f"{Path(command[-1]).stem}.pdf")
        return SimpleNamespace(returncode=0, stdout="", stderr="")

    monkeypatch.setattr(converter_module.subprocess, "run", fake_subprocess_run)

    class FakePDFToWord:
        def __init__(self, source):
            pass

        def convert(self, output):
            write_output(Path(output))

        def close(self):
            pass

    monkeypatch.setattr(converter_module, "Converter", FakePDFToWord, raising=False)

    class FakeFrame:
        def __init__(self, *args, **kwargs):
            pass

        def to_excel(self, writer, **kwargs):
            write_output(writer.path)

    def fake_writer(path, **kwargs):
        return nullcontext(SimpleNamespace(path=Path(path)))

    monkeypatch.setitem(
        sys.modules,
        "pandas",
        SimpleNamespace(DataFrame=FakeFrame, ExcelWriter=fake_writer),
    )
    monkeypatch.setattr(
        converter_module,
        "tabula",
        SimpleNamespace(read_pdf=lambda *args, **kwargs: [FakeFrame()]),
        raising=False,
    )
    monkeypatch.setattr(
        converter_module,
        "pdfplumber",
        SimpleNamespace(
            open=lambda _: nullcontext(
                SimpleNamespace(
                    pages=[
                        SimpleNamespace(extract_tables=lambda: [[["Name"], ["Alice"]]])
                    ]
                )
            )
        ),
        raising=False,
    )
    entrypoints = {
        "docx2pdf": converter._office_to_pdf_docx2pdf,
        "win32com": converter._office_to_pdf_win32com,
        "libreoffice": lambda source: converter._office_to_pdf_libreoffice(source, 5),
        "word": converter._pdf_to_word_sync,
        "tabula": converter._pdf_to_excel_tabula,
        "pdfplumber": converter._pdf_to_excel_pdfplumber,
    }
    return converter, input_path, extension, entrypoints[backend]


def _assert_distinct_concurrent_outputs(convert, source):
    with ThreadPoolExecutor(max_workers=2) as executor:
        outputs = list(executor.map(convert, [source, source]))
    assert all(output is not None for output in outputs)
    assert len(set(outputs)) == 2
    assert len({output.read_text(encoding="utf-8") for output in outputs}) == 2
    assert set(source.parent.iterdir()) == {source, *outputs}


@pytest.mark.parametrize("backend", _BACKENDS)
def test_conversion_preserves_existing_same_stem_file(monkeypatch, tmp_path, backend):
    rendered_paths = []

    def write_output(path):
        rendered_paths.append(path)
        path.write_bytes(b"new conversion")

    _, source, extension, convert = _fake_backend(
        monkeypatch, tmp_path, backend, write_output
    )
    existing = tmp_path / f"report{extension}"
    existing.write_bytes(b"existing output must survive")
    result = convert(source)
    assert result is not None and result != existing
    assert result.read_bytes() == b"new conversion"
    assert existing.read_bytes() == b"existing output must survive"
    assert source.read_bytes() == b"source must remain unchanged"
    assert all(path.parent != tmp_path for path in rendered_paths)
    assert set(tmp_path.iterdir()) == {source, existing, result}


@pytest.mark.parametrize("backend", _BACKENDS)
@pytest.mark.parametrize("failure", ["missing", "empty", "partial_error"])
def test_failed_conversion_never_reuses_old_output_or_leaves_files(
    monkeypatch, tmp_path, backend, failure
):
    def write_output(path):
        if failure == "empty":
            path.touch()
        if failure == "partial_error":
            path.write_bytes(b"incomplete conversion")
            raise RuntimeError("backend conversion failed")

    _, source, extension, convert = _fake_backend(
        monkeypatch, tmp_path, backend, write_output
    )
    existing = tmp_path / f"report{extension}"
    existing.write_bytes(b"old valid output")
    if failure == "partial_error" and backend in {"tabula", "pdfplumber"}:
        with pytest.raises(RuntimeError, match="backend conversion failed"):
            convert(source)
    else:
        assert convert(source) is None
    assert existing.read_bytes() == b"old valid output"
    assert source.read_bytes() == b"source must remain unchanged"
    assert set(tmp_path.iterdir()) == {source, existing}


@pytest.mark.parametrize("backend", _BACKENDS)
def test_concurrent_same_name_conversions_publish_distinct_outputs(
    monkeypatch, tmp_path, backend
):
    barrier = Barrier(2)

    def write_output(path):
        barrier.wait(timeout=5)
        path.write_text(path.parent.name, encoding="utf-8")

    _, source, _, convert = _fake_backend(monkeypatch, tmp_path, backend, write_output)
    _assert_distinct_concurrent_outputs(convert, source)


@pytest.mark.parametrize("source_suffix", [".xlsx", ".pptx"])
def test_docx2pdf_fallback_preserves_existing_pdf(monkeypatch, tmp_path, source_suffix):
    converter, source, _, _ = _fake_backend(
        monkeypatch,
        tmp_path,
        "docx2pdf",
        lambda path: path.write_bytes(b"fallback PDF"),
    )
    fallback_source = source.with_suffix(source_suffix)
    source.rename(fallback_source)
    existing = tmp_path / "report.pdf"
    existing.write_bytes(b"old PDF")
    result = converter._office_to_pdf_docx2pdf(fallback_source)
    assert result != existing and result.read_bytes() == b"fallback PDF"
    assert existing.read_bytes() == b"old PDF"
    assert set(tmp_path.iterdir()) == {fallback_source, existing, result}


def test_output_reservation_retries_when_concurrent_candidates_collide(
    monkeypatch, tmp_path
):
    converter, source, extension, convert = _fake_backend(
        monkeypatch,
        tmp_path,
        "word",
        lambda path: path.write_text(path.parent.name, encoding="utf-8"),
    )
    suggest_name = converter.get_unique_filename
    barrier = Barrier(2)
    lock = Lock()
    calls = 0

    def same_first_candidate(base_name, suffix):
        nonlocal calls
        with lock:
            calls += 1
            call = calls
        if call <= 2:
            barrier.wait(timeout=5)
            return tmp_path / f"report{extension}"
        return suggest_name(base_name, suffix)

    monkeypatch.setattr(converter, "get_unique_filename", same_first_candidate)
    _assert_distinct_concurrent_outputs(convert, source)
    assert calls == 3


def test_publish_failure_cleans_reserved_output(monkeypatch, tmp_path):
    _, source, extension, convert = _fake_backend(
        monkeypatch, tmp_path, "word", lambda path: path.write_bytes(b"converted")
    )
    existing = tmp_path / f"report{extension}"
    existing.write_bytes(b"old output")

    def fail_replace(self, target):
        raise OSError("publish failed")

    monkeypatch.setattr(Path, "replace", fail_replace)
    assert convert(source) is None
    assert existing.read_bytes() == b"old output"
    assert set(tmp_path.iterdir()) == {source, existing}


def test_real_pdf_to_word_keeps_existing_docx(tmp_path):
    fitz = pytest.importorskip("fitz")
    pdf2docx = pytest.importorskip("pdf2docx")
    from docx import Document

    assert converter_module.Converter is pdf2docx.Converter
    source = tmp_path / "report.pdf"
    with fitz.open() as pdf:
        page = pdf.new_page()
        page.insert_text((72, 72), "Fresh PDF conversion content")
        pdf.save(source)
    existing = tmp_path / "report.docx"
    document = Document()
    document.add_paragraph("Original Word source must survive")
    document.save(existing)
    original_bytes = existing.read_bytes()
    source_bytes = source.read_bytes()
    converter = PDFConverter.__new__(PDFConverter)
    converter.data_path = tmp_path
    output = converter._pdf_to_word_sync(source)
    assert output is not None and output != existing
    converted = Document(output)
    assert "Fresh PDF conversion content" in "\n".join(
        paragraph.text for paragraph in converted.paragraphs
    )
    assert existing.read_bytes() == original_bytes
    assert source.read_bytes() == source_bytes
    assert set(tmp_path.iterdir()) == {source, existing, output}
