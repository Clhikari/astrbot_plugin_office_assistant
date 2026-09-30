import pytest
from tests._docx_test_helpers import _write_png

from astrbot_plugin_office_assistant.constants import OfficeType
from astrbot_plugin_office_assistant.domain.document.contracts import (
    AddBlocksRequest,
    CreateDocumentRequest,
    ExportDocumentRequest,
)
from astrbot_plugin_office_assistant.domain.document.export_pipeline import (
    export_document_via_pipeline,
)
from astrbot_plugin_office_assistant.domain.document.render_backends import (
    NodeDocumentRenderBackend,
)
from astrbot_plugin_office_assistant.domain.document.session_store import (
    DocumentSessionStore,
)
from astrbot_plugin_office_assistant.services.word_read_service import WordReadService
from astrbot_plugin_office_assistant.services.workspace_service import WorkspaceService
from astrbot_plugin_office_assistant.utils import extract_ppt_text, extract_word_content


@pytest.mark.asyncio
@pytest.mark.parametrize("document_format", ["word", "ppt"])
async def test_real_node_exported_table_is_visible_to_project_reader(
    tmp_path, document_format
):
    docx = pytest.importorskip("docx")
    pptx = pytest.importorskip("pptx")
    backend = NodeDocumentRenderBackend()
    if not backend.is_available():
        pytest.skip("Build the Node renderer before running this integration test")
    store = DocumentSessionStore(workspace_dir=tmp_path)
    document = store.create_document(
        CreateDocumentRequest(title="Real roundtrip", format=document_format)
    )
    store.add_blocks(
        AddBlocksRequest(
            document_id=document.document_id,
            blocks=[
                {
                    "type": "table" if document_format == "word" else "table_slide",
                    "headers": ["项目", "数量", "标记"],
                    "rows": [["松针", "17", "CELL-A17"], ["银杏", "29", "CELL-B29"]],
                }
            ],
        )
    )
    _, output = await export_document_via_pipeline(
        store=store,
        render_backends=[backend],
        request=ExportDocumentRequest(document_id=document.document_id),
        source="real-table-reader-regression",
    )
    workspace = WorkspaceService(
        plugin_data_path=tmp_path,
        executor=None,
        office_libs={"docx": docx, "pptx": pptx},
        max_file_size=1024 * 1024,
    )
    if document_format == "word":
        reader = WordReadService(
            workspace_service=workspace, enable_docx_image_review=False
        )
        text = "\n".join(
            [
                result
                async for result in reader.iter_word_results(
                    output, output.name, ".docx", output.stat().st_size
                )
                if isinstance(result, str)
            ]
        )
    else:
        text = workspace.extract_office_text(output, OfficeType.POWERPOINT)
    assert text
    assert "项目\t数量\t标记" in text
    assert "松针\t17\tCELL-A17" in text
    assert "银杏\t29\tCELL-B29" in text


def test_ppt_reader_includes_tables_and_text_inside_nested_groups(tmp_path):
    pptx = pytest.importorskip("pptx")
    from pptx.util import Inches

    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(0, 0, Inches(2), Inches(1)).text = "Before group"
    table = slide.shapes.add_table(2, 2, 0, 0, Inches(4), Inches(2))
    table.table.cell(0, 0).text = "Name"
    table.table.cell(0, 1).text = "Value"
    table.table.cell(1, 0).text = "Nested-table-marker"
    table.table.cell(1, 1).text = "43"
    inner = slide.shapes.add_group_shape([table])
    outer = slide.shapes.add_group_shape([inner])
    outer.shapes.add_textbox(0, 0, Inches(2), Inches(1)).text = "Inside group"
    slide.shapes.add_textbox(0, 0, Inches(2), Inches(1)).text = "After group"
    output = tmp_path / "nested.pptx"
    presentation.save(output)
    assert extract_ppt_text(output) == (
        "Before group\nName\tValue\nNested-table-marker\t43\nInside group\nAfter group"
    )


@pytest.mark.parametrize("include_images", [False, True])
def test_word_table_reader_preserves_nested_content_and_inline_image_order(
    tmp_path, include_images
):
    docx = pytest.importorskip("docx")
    image = tmp_path / "pixel.png"
    _write_png(image, width=1, height=1)
    document = docx.Document()
    document.add_paragraph("Before table")
    table = document.add_table(rows=1, cols=2)
    cell = table.cell(0, 0)
    cell.paragraphs[0].add_run("Before cell image")
    cell.paragraphs[0].add_run().add_picture(str(image))
    cell.paragraphs[0].add_run("After cell image")
    nested = cell.add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "Nested-cell-marker"
    table.cell(0, 1).text = "Second cell"
    document.add_paragraph("After table")
    output = tmp_path / "nested.docx"
    document.save(output)
    result = extract_word_content(output, tmp_path, include_images=include_images)
    assert result is not None and result.text is not None
    markers = [
        "Before table",
        "Before cell image",
        "After cell image",
        "Nested-cell-marker",
        "Second cell",
        "After table",
    ]
    positions = [result.text.index(marker) for marker in markers]
    assert positions == sorted(positions)
    images = [item for item in result.items if item.type == "image"]
    assert len(images) == int(include_images)
    if include_images:
        image_index = next(
            i for i, item in enumerate(result.items) if item.type == "image"
        )
        assert "Before cell image" in result.items[image_index - 1].text
        assert "After cell image" in result.items[image_index + 1].text
