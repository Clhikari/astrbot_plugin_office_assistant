import json
import asyncio
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
from openpyxl import load_workbook

from astrbot_plugin_office_assistant.agent_tools import (
    build_document_toolset,
    build_workbook_toolset,
)


def context(user="alice", origin="private", platform="audit"):
    event = SimpleNamespace(
        get_sender_id=lambda: user,
        get_platform_id=lambda: platform,
        unified_msg_origin=origin,
        get_extra=lambda key, default=None: default,
    )
    return SimpleNamespace(context=SimpleNamespace(event=event))


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "other", [context("bob"), context(origin="other"), context(platform="other"), None]
)
@pytest.mark.parametrize(
    "kind,tool_name,arguments",
    [
        ("workbook", "write_rows", {"sheet": "Data", "rows": [["tampered"]]}),
        ("workbook", "export_workbook", {}),
        (
            "document",
            "add_blocks",
            {"blocks": [{"type": "paragraph", "text": "tampered"}]},
        ),
        (
            "document",
            "add_slides",
            {"slides": [{"type": "title_slide", "title": "tampered"}]},
        ),
        ("document", "finalize_document", {}),
        ("document", "export_document", {}),
    ],
)
async def test_owner_enforced_before_mutation_and_delivery(
    workspace_root, kind, tool_name, arguments, other
):
    delivery = AsyncMock()
    build = build_document_toolset if kind == "document" else build_workbook_toolset
    toolset = build(workspace_dir=workspace_root, after_export=delivery)
    tools = {tool.name: tool for tool in toolset.tools}

    async def call(name, actor, **kwargs):
        return json.loads(await tools[name].call(actor, **kwargs))

    created = await call(
        f"create_{kind}",
        context(),
        session_id="spoofed",
        title="private",
        **({"filename": "private.xlsx"} if kind == "workbook" else {}),
    )
    identifier = {f"{kind}_id": created[kind][f"{kind}_id"]}
    seed_name, seed_args = (
        ("add_blocks", {"blocks": [{"type": "paragraph", "text": "private"}]})
        if kind == "document"
        else ("write_rows", {"sheet": "Data", "rows": [["private", 12345]]})
    )
    assert (await call(seed_name, context(), **identifier, **seed_args))["success"]
    if tool_name == "export_document":
        assert (await call("finalize_document", context(), **identifier))["success"]
    store = getattr(toolset, f"{kind}_store")
    draft = getattr(store, f"require_{kind}")(*identifier.values())
    before = draft.model_dump()

    denied = await call(tool_name, other, **identifier, **arguments)
    assert denied["success"] is False
    assert "权限" in denied["message"]
    assert draft.model_dump() == before
    delivery.assert_not_awaited()
    if kind == "workbook":
        assert draft.worksheets[0].rows == [["private", 12345]]
        if tool_name == "export_workbook":
            assert await tools[tool_name].call(context(), **identifier) is None
            delivery.assert_awaited_once()


@pytest.mark.asyncio
async def test_invalid_event_cannot_create_an_unowned_workbook(workspace_root):
    toolset = build_workbook_toolset(workspace_dir=workspace_root)
    result = json.loads(
        await toolset.tools[0].call(
            SimpleNamespace(context=SimpleNamespace(event=None))
        )
    )
    assert result["success"] is False


@pytest.mark.asyncio
async def test_owned_workbooks_with_same_filename_keep_their_content_during_delivery(
    workspace_root,
):
    alice_waiting = asyncio.Event()
    bob_delivered = asyncio.Event()
    received = {}
    paths = {}

    async def deliver(tool_context, output_path):
        sender = tool_context.context.event.get_sender_id()
        paths[sender] = output_path
        if sender == "alice":
            alice_waiting.set()
            await bob_delivered.wait()
        with Path(output_path).open("rb") as stream:
            workbook = load_workbook(stream)
            received[sender] = workbook.active["A1"].value
            workbook.close()
        if sender == "bob":
            bob_delivered.set()

    toolset = build_workbook_toolset(workspace_dir=workspace_root, after_export=deliver)
    tools = {tool.name: tool for tool in toolset.tools}
    identifiers = {}
    for sender in ("alice", "bob"):
        created = json.loads(await tools["create_workbook"].call(context(sender)))
        identifiers[sender] = created["workbook"]["workbook_id"]
        await tools["write_rows"].call(
            context(sender),
            workbook_id=identifiers[sender],
            sheet="Data",
            rows=[[sender]],
        )
    alice_task = asyncio.create_task(
        tools["export_workbook"].call(context(), workbook_id=identifiers["alice"])
    )
    try:
        await asyncio.wait_for(alice_waiting.wait(), timeout=5)
        await tools["export_workbook"].call(
            context("bob"), workbook_id=identifiers["bob"]
        )
        await asyncio.wait_for(alice_task, timeout=5)
    finally:
        bob_delivered.set()
        await alice_task
    assert received == {"alice": "alice", "bob": "bob"}
    assert paths["alice"] != paths["bob"]
    assert all(Path(path).name == "workbook.xlsx" for path in paths.values())
